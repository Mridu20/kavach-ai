"""
Comprehensive Unit & Integration Test Suite for Person 1 (Lead Agentic AI).
"""

import unittest
from backend.agent.state import AgentState, HumanDecision, TaskCategory
from backend.agent.planner import AgentPlanner, TaskClassifier
from backend.agent.tools_registry import ToolRegistry
from backend.agent.orchestrator import AgentOrchestrator
from backend.agent.verifier import SelfVerifier
from backend.agent.human_approval import HumanApprovalManager


class TestAgentFramework(unittest.TestCase):

    def test_agent_state_creation(self):
        state = AgentState(task_id="test_001", user_query="Inspect turbine report")
        self.assertEqual(state.task_id, "test_001")
        self.assertEqual(state.status, "INITIALIZED")
        self.assertEqual(len(state.trace.events), 0)

        state.add_trace_event("TEST_EVENT", "State initialized cleanly")
        self.assertEqual(len(state.trace.events), 1)
        self.assertEqual(state.trace.events[0].event_type, "TEST_EVENT")

    def test_task_classification_and_planning(self):
        # 1. Inspection Report
        category = TaskClassifier.classify("Inspect scanned inspection report", ["report.pdf"])
        self.assertEqual(category, TaskCategory.DOCUMENT_INSPECTION)

        state = AgentState(task_id="test_002", user_query="Inspect scanned report", input_files=["report.pdf"])
        cat, steps = AgentPlanner.create_plan(state)
        self.assertEqual(cat, TaskCategory.DOCUMENT_INSPECTION)
        self.assertEqual(len(steps), 6)
        self.assertEqual(steps[0].assigned_tool, "ocr_pdf_tool")

        # 2. SOP Query
        cat_sop, steps_sop = AgentPlanner.create_plan(AgentState(task_id="t2", user_query="What is the SOP policy?"))
        self.assertEqual(cat_sop, TaskCategory.SOP_RAG_QUERY)
        self.assertEqual(len(steps_sop), 2)

    def test_tool_registry_and_retry(self):
        registry = ToolRegistry()
        tools = registry.list_tools()
        self.assertGreaterEqual(len(tools), 7)

        # Execute valid tool
        rec = registry.execute_with_retry(tool_name="ocr_pdf_tool", step_id=1, params={"file_path": "sample.pdf"})
        self.assertTrue(rec.success)
        self.assertIn("extracted_text", rec.output)

        # Execute invalid tool
        rec_invalid = registry.execute_with_retry(tool_name="non_existent_tool", step_id=1, params={})
        self.assertFalse(rec_invalid.success)
        self.assertIn("not registered", rec_invalid.error_message)

    def test_orchestrator_execution(self):
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        self.assertEqual(state.status, "PLANNED")
        self.assertEqual(len(state.plan), 6)

        orchestrator.run_all_steps(state)
        self.assertEqual(state.status, "VERIFYING")
        self.assertEqual(len(state.tool_calls), 6)
        self.assertGreaterEqual(len(state.retrieved_evidence), 1)
        self.assertIn("approval_note_docx", state.draft_deliverables)
        self.assertIn("action_tracker_xlsx", state.draft_deliverables)

    def test_self_verification(self):
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        orchestrator.run_all_steps(state)

        result = SelfVerifier.verify(state)
        self.assertTrue(result.verified)
        self.assertTrue(result.zero_external_calls)
        self.assertEqual(state.status, "AWAITING_APPROVAL")

    def test_human_approval_gate(self):
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        orchestrator.run_all_steps(state)
        SelfVerifier.verify(state)

        # Approve
        updated_state = HumanApprovalManager.submit_decision(
            state=state,
            decision=HumanDecision.APPROVED,
            reviewer="Lead Inspector",
            comments="All verified.",
        )
        self.assertEqual(updated_state.status, "COMPLETED")
        self.assertEqual(updated_state.approval.reviewer, "Lead Inspector")
        self.assertEqual(updated_state.approval.status, HumanDecision.APPROVED)

    def test_model_router_output_reaches_findings(self):
        """Regression test for ROOT CAUSE: model_router_tool output was silently dropped.
        
        Asserts that model_router_tool writes synthesized_analysis into state.findings,
        and that it is non-empty and substantive.
        """
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        orchestrator.run_all_steps(state)

        # model_used must always be set
        self.assertIn("model_used", state.findings,
            "model_router_tool output should be written to state.findings['model_used']")

        # synthesized_analysis MUST reach state.findings and be non-empty
        self.assertIn("synthesized_analysis", state.findings,
            "synthesized_analysis must be present in state.findings when model_router_tool runs")
        synthesis = state.findings["synthesized_analysis"]
        self.assertIsInstance(synthesis, str)
        self.assertGreater(len(synthesis.strip()), 50,
            "state.findings['synthesized_analysis'] must contain substantive synthesized content, not an empty string")

    def test_xlsx_items_derived_from_findings(self):
        """Regression test: XLSX items were hardcoded to one fake task instead of real findings."""
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        orchestrator.run_all_steps(state)

        # Find the xlsx tool call and check its input_params
        xlsx_calls = [tc for tc in state.tool_calls if tc.tool_name == "generate_xlsx_tool"]
        self.assertEqual(len(xlsx_calls), 1)
        items = xlsx_calls[0].input_params.get("items", [])
        self.assertGreaterEqual(len(items), 2, "XLSX should have multiple items derived from real findings and evidence")

        # Confirm items have valid priority values
        for item in items:
            self.assertIn(item.get("priority"), ["LOW", "MEDIUM", "HIGH", "CRITICAL"])
            self.assertGreater(len(item.get("task", "")), 5)

        # Confirm XLSX on disk has real content and rows
        import os
        from openpyxl import load_workbook
        xlsx_file = state.draft_deliverables.get("action_tracker_xlsx", "")
        self.assertTrue(xlsx_file, "Action tracker deliverable path must be set")
        full_xlsx_path = os.path.join("backend", "storage", "outputs", xlsx_file)
        self.assertTrue(os.path.exists(full_xlsx_path))
        wb = load_workbook(full_xlsx_path)
        ws = wb.active
        self.assertGreaterEqual(ws.max_row, 3, "XLSX should have header row plus multiple data rows")

    def test_docx_content_quality(self):
        """Content-quality test: DOCX must contain synthesized analysis and citations, not just a title."""
        import os
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="Inspect turbine report", input_files=["turbine.pdf"])
        orchestrator.run_all_steps(state)

        # Verify the DOCX file was created
        docx_path = state.draft_deliverables.get("approval_note_docx", "")
        self.assertTrue(docx_path, "DOCX deliverable path should be set")

        full_path = os.path.join("backend", "storage", "outputs", docx_path)
        self.assertTrue(os.path.exists(full_path), f"DOCX file should exist at {full_path}")

        # Open and verify actual content quality
        from docx import Document
        doc = Document(full_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

        self.assertIn("Executive Summary", headings, "DOCX must contain 'Executive Summary' heading")
        self.assertIn("Detailed Findings", headings, "DOCX must contain 'Detailed Findings' heading")
        self.assertIn("SOP Evidence & Citations", headings, "DOCX must contain 'SOP Evidence & Citations' heading")
        self.assertIn(f"Task Reference: {state.task_id}", full_text, "DOCX must contain explicit task reference ID")

        # Verify that state.findings["synthesized_analysis"] is actually consumed in the DOCX text
        synthesized = state.findings.get("synthesized_analysis", "")
        self.assertTrue(synthesized, "Synthesized analysis must be present in state.findings")
        # Check that key words or sentences from synthesized analysis appear in docx full_text
        first_meaningful_line = next((l.strip().lstrip("#*-0123456789. ") for l in synthesized.splitlines() if len(l.strip()) > 20), "")
        if first_meaningful_line:
            clean_check = first_meaningful_line.replace("**", "")[:40]
            self.assertIn(clean_check, full_text, "DOCX Executive Summary must include content from synthesized_analysis")

        # Verify that SOP evidence citations from state.retrieved_evidence appear in the DOCX
        self.assertGreater(len(state.retrieved_evidence), 0, "State should have retrieved SOP evidence")
        first_ev = state.retrieved_evidence[0]
        self.assertIn(first_ev.source_doc, full_text, "DOCX must include the SOP citation document name")

    def test_conversational_input_hi_produces_clean_greeting(self):
        """Conversational guard test: 'hi' must produce a clean greeting, NOT an invented risk template."""
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(query="hi")
        orchestrator.run_all_steps(state)

        # Must be classified as GENERAL_REASONING
        self.assertEqual(state.category, TaskCategory.GENERAL_REASONING)

        # Model router call must have task_type='general'
        model_calls = [tc for tc in state.tool_calls if tc.tool_name == "model_router_tool"]
        self.assertEqual(len(model_calls), 1)
        self.assertEqual(model_calls[0].input_params.get("task_type"), "general")

        # Response must be a clean greeting and intro to KAVACH AI
        response = state.text_response or state.findings.get("synthesized_analysis", "")
        self.assertTrue(response, "Agent must return a text response for 'hi'")
        self.assertIn("KAVACH AI", response, "Greeting should introduce KAVACH AI")
        self.assertTrue(any(w in response.lower() for w in ["hello", "assist", "welcome", "sovereign"]),
            "Response should be a welcoming greeting")

        # Response must NOT contain placeholder brackets or fake risk template headers
        self.assertNotIn("[insert specific", response.lower())
        self.assertNotIn("[insert areas", response.lower())
        self.assertNotIn("sample risk assessment", response.lower())

    def test_rag_search_fallback_flagged_honestly(self):
        """Honest fallback test: when RAG fails/falls back, it is explicitly tagged."""
        from backend.agent.tools_registry import RAGSearchTool
        tool = RAGSearchTool()

        # Call with invalid query to trigger store search or exception fallback
        res = tool.run(query="___non_existent_trigger___")
        self.assertIn("results", res)

        # If fallback was activated, ensure fallback flags are set honestly
        if res.get("fallback_data"):
            self.assertTrue(res.get("FALLBACK_DATA"), "FALLBACK_DATA flag must be True")
            self.assertTrue(res.get("is_fallback"), "is_fallback flag must be True")
            first_result = res["results"][0]
            self.assertTrue(first_result.get("fallback") or first_result.get("is_fallback"),
                "Individual fallback items must be tagged with fallback=True")


    def test_streaming_agent_endpoint(self):
        """Test FastAPI SSE streaming route returns real-time data events."""
        from fastapi.testclient import TestClient
        from backend.main import app

        client = TestClient(app)
        response = client.post(
            "/api/agent/run-stream",
            json={"user_query": "Inspect turbine report", "input_files": ["turbine.pdf"]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/event-stream", response.headers.get("content-type", ""))

        content = response.text
        self.assertIn("data: ", content)
        self.assertIn('"type": "INIT"', content)
        self.assertIn('"type": "COMPLETE"', content)

    def test_cancellation_and_fields(self):
        """Test cancellation state handling and calculation/multi-doc detection."""
        orchestrator = AgentOrchestrator()
        state = orchestrator.create_task(
            query="Calculate derated MAWP for pressure vessel with measured thickness 0.285 in",
            input_files=["baseline.pdf", "current.pdf"],
        )
        # Verify initial state
        self.assertFalse(state.cancellation_requested)
        # Execute 1 step
        orchestrator.execute_next_step(state)
        # Request cancel
        state.request_cancel()
        self.assertTrue(state.cancellation_requested)
        self.assertEqual(state.status, "CANCELLED")

        # Verify execute_next_step returns CANCELLED
        state2 = orchestrator.execute_next_step(state)
        self.assertEqual(state2.status, "CANCELLED")

        # Test full run with calculation populates calculation_details and multi_doc_comparison
        state_calc = orchestrator.create_task(
            query="Calculate derated MAWP for pressure vessel",
            input_files=["baseline.pdf", "current.pdf"],
        )
        orchestrator.run_all_steps(state_calc)
        self.assertIsNotNone(state_calc.text_response)
        self.assertIsNotNone(state_calc.calculation_details)
        self.assertIn("steps", state_calc.calculation_details)
        self.assertIsNotNone(state_calc.multi_doc_comparison)
        self.assertEqual(state_calc.multi_doc_comparison["doc1"], "baseline.pdf")


if __name__ == "__main__":
    unittest.main()
